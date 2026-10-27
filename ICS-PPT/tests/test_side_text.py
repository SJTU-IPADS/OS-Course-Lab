"""Slide text set under a side image: `p.slide(...).under_side_image()`.

The image column, on either side, then holds the picture at its top and the
text below it; the other blocks keep the text column.
"""

import base64
import tempfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from lecturekit import Lecture, model, serialize
from lecturekit.renderers.pptx import PptxRenderer
from lecturekit.renderers.pptx.layout import Layout, px
from lecturekit.renderers.viewer import render_marp_page


_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+M9Q"
    "DwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


def _lecture(side="left", *, image=True):
    lecture = Lecture(id="lec", title="T")

    def body(p):
        p.title("Title")
        if image:
            p.side_image("assets/pic.png", width="40%", alt="contain", side=side)
        p.slide("- under the picture").under_side_image()
        p.code("c", "int beside;")
    lecture.page("pg", body=body)
    return lecture


def test_the_handle_marks_the_slide_block():
    page = _lecture().build().children[0]
    marked = [block.kind for block in page.blocks if block.in_side]
    assert marked == ["slide"]
    blocks = serialize.lecture_to_dict(_lecture().build())["children"][0]["blocks"]
    assert [block["in_side"] for block in blocks] == [False, True, False]


def test_only_a_slide_block_may_be_set_under_the_image():
    lecture = Lecture(id="lec", title="T")

    def body(p):
        p.title("Title")
        p.side_image("assets/pic.png", side="left")
        p.code("c", "int x;").under_side_image()
    lecture.page("pg", body=body)
    with pytest.raises(model.ValidationError, match="slide block"):
        lecture.build()


def test_it_needs_a_side_image():
    with pytest.raises(model.ValidationError, match="side_image"):
        _lecture(image=False).build()
    _lecture("right").build()                   # either side holds the text


def test_the_viewer_draws_the_text_in_the_image_column():
    md = render_marp_page(_lecture().build().children[0], reveal=True)
    lines = md.split("\n")
    start = lines.index('<div class="lk-side lk-side-stack">')
    end = lines.index("</div>", start)
    # the images come first, on the line that keeps the HTML block open
    assert lines[start + 1] == (
        '<div class="lk-side-images">'
        '<img src="assets/pic.png" alt="" class="lk-side-fit"></div>'
    )
    # the text is markdown of its own, between blank lines, inside the column
    assert lines[start + 2] == ""
    assert "- under the picture" in lines[start + 3:end]
    assert md.count("under the picture") == 1
    # the column is not a reveal step; the code block beside it is the only one
    assert md.count('class="reveal-block"') == 1
    assert md.index("</div>", md.index("lk-side-stack")) < md.index("int beside;")


def test_a_column_without_text_is_drawn_as_before():
    lecture = Lecture(id="lec", title="T")

    def body(p):
        p.title("Title")
        p.side_image("assets/pic.png", width="40%", alt="contain", side="left")
        p.slide("beside the picture")
    lecture.page("pg", body=body)
    md = render_marp_page(lecture.build().children[0])
    assert '<figure class="lk-side">' in md
    assert "lk-side-stack" not in md


def test_powerpoint_stacks_the_picture_and_the_text_in_the_column():
    src = Path(tempfile.mkdtemp())
    (src / "assets").mkdir()
    (src / "assets" / "pic.png").write_bytes(_PNG)
    out = PptxRenderer(asset_root=src).render(_lecture().build(), Path(tempfile.mkdtemp()))
    slide = Presentation(str(out)).slides[0]
    layout = Layout.from_ratio("16:9")
    column = round(layout.width * 0.4)
    pic = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    under = next(s for s in slide.shapes if s.has_text_frame
                 and s.text_frame.text.endswith("under the picture"))
    beside = next(s for s in slide.shapes if s.has_text_frame
                  and "int beside;" in s.text_frame.text)
    # the picture is as wide as the column, which keeps the slide's margins
    assert pic.left == layout.content_left
    assert pic.left + pic.width == column + px(40)
    # the text is in the same column, below the picture
    assert under.left == pic.left
    assert under.top >= pic.top + pic.height
    assert under.left + under.width <= pic.left + pic.width
    # the other block keeps the text column and starts under the title
    assert beside.left >= column + layout.content_left
    assert beside.top < under.top


def test_the_viewer_stacks_a_right_column_the_same_way():
    md = render_marp_page(_lecture("right").build().children[0])
    assert md.startswith("<!-- _class: lk-side-right -->")
    assert '<div class="lk-side lk-side-stack">' in md
    assert md.count("under the picture") == 1
    assert md.index("</div>", md.index("lk-side-stack")) < md.index("int beside;")


def test_powerpoint_stacks_a_right_column_against_the_right_margin():
    src = Path(tempfile.mkdtemp())
    (src / "assets").mkdir()
    (src / "assets" / "pic.png").write_bytes(_PNG)
    out = PptxRenderer(asset_root=src).render(
        _lecture("right").build(), Path(tempfile.mkdtemp()))
    slide = Presentation(str(out)).slides[0]
    layout = Layout.from_ratio("16:9")
    column = round(layout.width * 0.4)
    pic = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    under = next(s for s in slide.shapes if s.has_text_frame
                 and s.text_frame.text.endswith("under the picture"))
    beside = next(s for s in slide.shapes if s.has_text_frame
                  and "int beside;" in s.text_frame.text)
    # the mirror of the left column: the outer edge at the side padding, the
    # inner edge 40px into the text column's share
    assert pic.left + pic.width == layout.width - layout.content_left
    assert pic.left == layout.width - column - px(40)
    assert under.left == pic.left
    assert under.top >= pic.top + pic.height
    assert beside.left == layout.content_left
    assert beside.left + beside.width <= layout.content_left + layout.content_width - column
