"""Both columns of a side_image page in the vertical middle:
`p.side_image(..., valign="middle")`.

The text column's blocks are centred in the band from under the title to the
slide's bottom padding; the image column (a fitted picture, or a stack with
text under it) is centred in the same band, on either side.
"""

import base64
import tempfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from lecturekit import Lecture, model
from lecturekit.renderers.pptx import PptxRenderer
from lecturekit.renderers.pptx.layout import Layout
from lecturekit.renderers.viewer import render_marp_page


_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+M9Q"
    "DwADhgGAWjR9awAAAABJRU5ErkJggg=="
)

_ROWS = [["a", "1"], ["b", "2"], ["c", "3"]]


def _lecture(valign="middle", *, side="left", under=True, gap=False):
    lecture = Lecture(id="lec", title="T")

    def body(p):
        p.title("Title")
        if gap:
            p.gap()
        p.side_image("assets/pic.png", width="40%", alt="contain",
                     side=side, valign=valign)
        if under:
            p.slide("- under the picture").under_side_image()
        p.table(_ROWS, headers=["k", "v"])
        p.code("c", "int beside;")
    lecture.page("pg", body=body)
    return lecture


def _page(**kw):
    return _lecture(**kw).build().children[0]


def _render(lecture):
    src = Path(tempfile.mkdtemp())
    (src / "assets").mkdir()
    (src / "assets" / "pic.png").write_bytes(_PNG)
    out = PptxRenderer(asset_root=src).render(lecture.build(), Path(tempfile.mkdtemp()))
    return Presentation(str(out)).slides[0]


def _extent(shapes):
    return min(s.top for s in shapes), max(s.top + s.height for s in shapes)


def _column_shapes(slide):
    """The text column's shapes (the table and the code block) and, for a
    stack, the picture with the text under it; the band they are centred in
    starts under the title's rule."""
    shapes = list(slide.shapes)      # one proxy per shape: `in` compares them
    pic = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    rule = next(s for s in shapes if not s.has_text_frame and s not in pic)
    under = [s for s in shapes if s.has_text_frame
             and s.text_frame.text.endswith("under the picture")]
    column = [s for s in shapes
              if s not in pic and s not in under and s != rule
              and not (s.has_text_frame and s.text_frame.text == "Title")]
    return rule, pic + under, column


# --- the DSL and validation ---

def test_the_content_carries_the_alignment():
    page = _page()
    side = next(b for b in page.blocks if b.kind == "side_image")
    assert side.content["valign"] == "middle"
    side = next(b for b in _page(valign="top").blocks if b.kind == "side_image")
    assert side.content["valign"] == "top"


def test_a_block_without_the_key_is_at_the_top():
    block = model.Block(kind="side_image", content={"src": "x.png", "side": "left"})
    assert model.side_valign([block]) == "top"


def test_the_dsl_rejects_an_unknown_alignment():
    # "bottom" is the page's third valign (docs/layout.md); "center" names
    # the other axis.
    with pytest.raises(model.ValidationError, match="valign"):
        _lecture(valign="center").build()


def test_centring_and_a_page_gap_exclude_each_other():
    with pytest.raises(model.ValidationError, match="gap"):
        _lecture(gap=True).build()
    _lecture(valign="top", gap=True).build()        # the default keeps p.gap


def test_the_side_images_of_a_page_must_agree():
    lecture = Lecture(id="lec", title="T")

    def body(p):
        p.title("Title")
        p.side_image("a.png", side="left", valign="middle")
        p.side_image("b.png", side="left")
        p.slide("beside")
    lecture.page("pg", body=body)
    with pytest.raises(model.ValidationError, match="disagree"):
        lecture.build()


# --- the viewer ---

def test_the_viewer_wraps_the_text_column_in_two_spacers():
    md = render_marp_page(_page(), reveal=True)
    lines = md.split("\n")
    assert lines[0] == "<!-- _class: lk-side-left lk-valign-middle -->"
    spacer = '<div class="lk-valign-spacer" aria-hidden="true"></div>'
    assert md.count(spacer) == 2
    first = lines.index(spacer)
    last = lines.index(spacer, first + 1)
    # the first spacer follows the image column, the last one ends the page;
    # each is its own HTML block between blank lines
    stack = lines.index('<div class="lk-side lk-side-stack">')
    assert lines.index("</div>", stack) < first
    assert lines[first - 1] == "" and lines[first + 1] == ""
    assert lines[last - 1] == "" and last == len(lines) - 1
    column = lines[first:last]
    assert any("| a | 1 |" in line for line in column)
    assert any("int beside;" in line for line in column)
    # the spacers are not reveal steps: the table and the code block are
    assert md.count("reveal-block") == 2


def test_the_viewer_centres_a_right_column_page_the_same_way():
    md = render_marp_page(_page(side="right"))
    assert md.startswith("<!-- _class: lk-side-right lk-valign-middle -->")
    assert md.count('class="lk-valign-spacer"') == 2
    assert '<div class="lk-side lk-side-stack">' in md
    assert "![bg" not in md


def test_the_viewer_draws_the_default_as_before():
    md = render_marp_page(_page(valign="top"))
    assert "lk-valign" not in md
    assert md.startswith("<!-- _class: lk-side-left -->")


# --- PowerPoint ---

def test_powerpoint_centres_both_columns_on_one_line():
    rule, stack, column = _column_shapes(_render(_lecture()))
    layout = Layout.from_ratio("16:9")
    # the band starts where both groups start by default: under the title
    _, plain_stack, plain_column = _column_shapes(_render(_lecture(valign="top")))
    band_top = _extent(plain_stack)[0]
    assert band_top == _extent(plain_column)[0] > rule.top
    middle = (band_top + layout.content_bottom) / 2
    top, bottom = _extent(column)
    assert abs((top + bottom) / 2 - middle) <= 1
    top, bottom = _extent(stack)
    assert abs((top + bottom) / 2 - middle) <= 1
    # the geometry of each group is the renderer's own: the text stays under
    # the picture, the code block under the table
    pic = next(s for s in stack if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    under = next(s for s in stack if s.has_text_frame)
    assert under.top >= pic.top + pic.height


@pytest.mark.parametrize("side", ["left", "right"])
def test_powerpoint_centres_both_columns_on_either_side(side):
    rule, stack, column = _column_shapes(_render(_lecture(side=side)))
    layout = Layout.from_ratio("16:9")
    _, plain_stack, _ = _column_shapes(_render(_lecture(side=side, valign="top")))
    middle = (_extent(plain_stack)[0] + layout.content_bottom) / 2
    for group in (column, stack):
        top, bottom = _extent(group)
        assert abs((top + bottom) / 2 - middle) <= 1
    # the stack keeps its side of the slide
    pic = next(s for s in stack if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    if side == "left":
        assert pic.left + pic.width < min(s.left for s in column)
    else:
        assert pic.left > max(s.left + s.width for s in column)


def test_powerpoint_centres_a_fitted_picture_in_the_same_band():
    slide = _render(_lecture(under=False))
    rule, pics, column = _column_shapes(slide)
    layout = Layout.from_ratio("16:9")
    pic = pics[0]
    top, bottom = _extent(column)
    assert abs((pic.top + pic.top + pic.height) / 2 - (top + bottom) / 2) <= 1
    assert pic.top + pic.height <= layout.content_bottom


def test_powerpoint_leaves_the_default_at_the_top():
    centred = _render(_lecture())
    plain = _render(_lecture(valign="top"))
    _, stack_c, column_c = _column_shapes(centred)
    _, stack_p, column_p = _column_shapes(plain)
    assert _extent(column_c)[0] > _extent(column_p)[0]
    assert _extent(stack_c)[0] > _extent(stack_p)[0]
    # moving the group keeps its height
    assert (_extent(column_c)[1] - _extent(column_c)[0]
            == _extent(column_p)[1] - _extent(column_p)[0])
