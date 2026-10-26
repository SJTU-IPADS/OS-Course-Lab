"""A small image beside slide text, centred against it:
`p.slide(...).image_right(..., valign="middle")`.

The image and the text sit side by side, each centred vertically against the
other; the default keeps the float the text wraps around.
"""

import base64
import tempfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from lecturekit import Lecture, model
from lecturekit.model import Block, check_block
from lecturekit.renderers.pptx import PptxRenderer
from lecturekit.renderers.viewer.blocks import render_block


_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+M9Q"
    "DwADhgGAWjR9awAAAABJRU5ErkJggg=="
)

_TEXT = "**idea**:\n- one\n- two\n- three\n- four"


def _lecture(valign="middle", *, text=_TEXT, height_px=60):
    lecture = Lecture(id="lec", title="T")

    def body(p):
        p.title("Title")
        p.slide(text).image_right("assets/pic.png", width_px=200,
                                  height_px=height_px, valign=valign)
    lecture.page("pg", body=body)
    return lecture


# --- the DSL and validation ---

def test_the_float_image_carries_the_alignment():
    block = _lecture().build().children[0].blocks[0]
    assert block.float_image["valign"] == "middle"


def test_the_dsl_rejects_an_unknown_alignment():
    with pytest.raises(model.ValidationError, match="valign"):
        _lecture("bottom").build()


def test_validation_rejects_an_unknown_alignment_in_the_tree():
    block = Block(kind="slide", content="x",
                  float_image={"src": "a.png", "valign": "bottom"})
    with pytest.raises(model.ValidationError, match="valign"):
        check_block(block, page_id="p")
    # a block written before the key existed is at the top
    check_block(Block(kind="slide", content="x", float_image={"src": "a.png"}),
                page_id="p")


# --- the viewer ---

def _middle_block(**extra):
    image = {"src": "pic.svg", "alt": "", "width": "456px", "height": None,
             "side": "right", "valign": "middle"}
    return Block(kind="slide", content="the idea", float_image=image, **extra)


def test_the_viewer_sets_text_and_image_side_by_side():
    lines = render_block(_middle_block())
    assert lines[:3] == ['<div class="lk-float lk-float-middle">',
                         '<div class="lk-float-text">', ""]
    # the text is markdown of its own, then the image follows it on the right
    assert lines[3] == "the idea"
    assert lines[4:6] == ["", "</div>"]
    assert lines[6].startswith("<img ") and "float" not in lines[6]
    assert "width: 456px" in lines[6]
    assert lines[7:] == ["</div>", ""]


def test_the_footnote_marker_stays_on_the_text():
    lines = render_block(_middle_block(footnotes=("n",)), footnote_numbers=(1,))
    assert "<sup" in lines[3]
    assert "<sup" not in lines[6]


def test_the_default_keeps_the_float():
    block = _middle_block()
    block = Block(kind="slide", content="the idea",
                  float_image={**block.float_image, "valign": "top"})
    lines = render_block(block)
    assert lines[0] == '<div class="lk-float">'
    assert "float: right" in lines[2]


# --- PowerPoint ---

def _render(lecture):
    src = Path(tempfile.mkdtemp())
    (src / "assets").mkdir()
    (src / "assets" / "pic.png").write_bytes(_PNG)
    out = PptxRenderer(asset_root=src).render(lecture.build(), Path(tempfile.mkdtemp()))
    slide = Presentation(str(out)).slides[0]
    pic = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    text = next(s for s in slide.shapes if s.has_text_frame
                and "one" in s.text_frame.text)
    return pic, text


def _middle(shape):
    return shape.top + shape.height / 2


def test_powerpoint_centres_a_short_picture_against_the_text():
    pic, text = _render(_lecture())
    assert pic.height < text.height
    assert abs(_middle(pic) - _middle(text)) <= 1
    top, _ = _render(_lecture("top"))
    assert top.top == text.top           # the default hugs the text's top


def test_powerpoint_centres_short_text_against_the_picture():
    pic, text = _render(_lecture(text="- one", height_px=200))
    assert text.height < pic.height
    assert abs(_middle(pic) - _middle(text)) <= 1
    _, plain = _render(_lecture("top", text="- one", height_px=200))
    assert text.top > plain.top
