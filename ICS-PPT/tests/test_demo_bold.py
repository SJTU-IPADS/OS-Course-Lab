import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from lecturekit import Lecture
from lecturekit.model import ValidationError
from lecturekit.renderers.latex.blocks import Ctx, emit_block
from lecturekit.renderers.pptx import PptxRenderer
from lecturekit.renderers.viewer.blocks import render_block

OUTPUT = """
scheme   rmse
Q4_0     0.100460
Q4_1     0.022288
"""


def demo_lecture(output=OUTPUT, **kw):
    lecture = Lecture(id="lec", title="L")

    def body(p):
        p.title("P")
        p.demo("compare", "./quant_compare w.bf16", output=output, **kw)

    lecture.page("p1", body=body)
    return lecture


def demo_block(**kw):
    blocks = demo_lecture(**kw).build().children[0].blocks
    return next(b for b in blocks if b.kind == "demo")


class BoldAuthoringTest(unittest.TestCase):
    def test_bold_is_stored_only_when_given(self):
        self.assertNotIn("bold", demo_block().content)
        self.assertEqual(demo_block(bold=[2, 3]).content["bold"], [2, 3])

    def test_numbers_count_the_output_with_its_blank_edges_stripped(self):
        demo_lecture(bold=[3]).build()  # line 3 is Q4_1, the last one

    def test_a_line_past_the_end_is_refused(self):
        with self.assertRaises(ValidationError) as ctx:
            demo_lecture(bold=[4]).build()
        self.assertIn("3 lines", str(ctx.exception))

    def test_zero_and_booleans_are_refused(self):
        for bad in (0, True, "2"):
            with self.subTest(bad=bad), self.assertRaises(ValidationError):
                demo_lecture(bold=[bad]).build()

    def test_a_blank_line_is_refused(self):
        with self.assertRaises(ValidationError) as ctx:
            demo_lecture(output="a\n\nb", bold=[2]).build()
        self.assertIn("blank", str(ctx.exception))

    def test_bold_needs_an_output(self):
        with self.assertRaises(ValidationError) as ctx:
            demo_lecture(output=None, bold=[1]).build()
        self.assertIn("needs an output", str(ctx.exception))


class BoldRenderTest(unittest.TestCase):
    def test_the_deck_wraps_the_named_lines(self):
        html = render_block(demo_block(bold=[2, 3]))[0]
        self.assertIn("<strong>Q4_0     0.100460</strong>", html)
        self.assertIn("<strong>Q4_1     0.022288</strong>", html)
        self.assertNotIn("<strong>scheme", html)

    def test_the_deck_marks_the_output_so_the_rest_can_recede(self):
        html = render_block(demo_block(bold=[2]))[0]
        self.assertIn('<code class="lk-demo-out" data-lk-demo-bold>', html)

    def test_the_deck_is_unchanged_without_bold(self):
        html = render_block(demo_block())[0]
        self.assertNotIn("<strong>", html)
        self.assertIn('<code class="lk-demo-out">\nscheme', html)

    def test_the_pptx_bolds_the_named_rows(self):
        out = PptxRenderer().render(demo_lecture(bold=[2]).build(),
                                    Path(tempfile.mkdtemp()))
        runs = {
            run.text: run.font.bold
            for shape in Presentation(str(out)).slides[0].shapes
            if shape.has_text_frame
            for para in shape.text_frame.paragraphs
            for run in para.runs
        }
        self.assertTrue(runs["Q4_0     0.100460"])
        self.assertFalse(runs["Q4_1     0.022288"])
        self.assertFalse(runs["$ ./quant_compare w.bf16"])

    def test_the_book_sets_named_lines_bold_and_the_rest_light(self):
        ctx = Ctx(lecture_id="lec", page_id="p1", slide_width=1280, assets=None)
        out = emit_block(demo_block(bold=[3]), ctx)
        self.assertIn(r"moredelim={[is][\bfseries]{(@}{@)}}", out)
        self.assertIn(r"moredelim={[is][\fontseries{l}\selectfont]{<@}{@>}}", out)
        self.assertIn("(@Q4_1     0.022288@)", out)
        self.assertIn("<@Q4_0     0.100460@>", out)
        self.assertIn("<@scheme   rmse@>", out)
        self.assertIn("$ ./quant_compare w.bf16\n", out)

    def test_the_book_picks_delimiters_the_output_lacks(self):
        ctx = Ctx(lecture_id="lec", page_id="p1", slide_width=1280, assets=None)
        out = emit_block(demo_block(output="f(@x)\ny", bold=[2]), ctx)
        self.assertIn("<@y@>", out)
        self.assertIn("[@f(@x)@]", out)


if __name__ == "__main__":
    unittest.main()
