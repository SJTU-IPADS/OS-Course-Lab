"""``p.flex("row", widths=[...])``: a row whose children take unequal shares."""

import pytest
from PIL import Image
from pptx import Presentation

from lecturekit import Lecture, model
from lecturekit.dsl import PageBuilder
from lecturekit.renderers.pptx import PptxRenderer
from lecturekit.renderers.pptx.layout import Layout, px
from lecturekit.renderers.viewer import render_marp_page
from lecturekit.serialize import lecture_to_dict


def _lecture(body):
    lecture = Lecture(id="widths", title="Widths")
    lecture.page("page", body=body)
    return lecture


def _page(body):
    return _lecture(body).build().children[0]


def _row(widths):
    def body(p):
        p.title("Title")
        with p.flex("row", widths=widths):
            p.slide("one", autobold=False)
            p.code("c", "two")
    return body


def _pptx(body, tmp_path):
    Image.new("RGB", (200, 100), "white").save(tmp_path / "image.png")
    out = PptxRenderer(asset_root=tmp_path).render(_lecture(body).build(), tmp_path)
    return Presentation(str(out)).slides[0]


def test_weights_are_stored_on_the_children_as_fractions():
    page = _page(_row([1, 3]))
    one, two = (b for b in page.blocks if b.kind in ("slide", "code"))
    assert (one.weight, two.weight) == (0.25, 0.75)
    assert page.blocks[0].kind == model.FLEX_KIND and page.blocks[0].weight is None
    assert lecture_to_dict(_lecture(_row([1, 3])).build())["children"][0]["blocks"][1]["weight"] == 0.25


def test_a_nested_container_counts_once_and_takes_its_weight():
    def body(p):
        p.title("Title")
        with p.flex("row", widths=[1, 1]):
            with p.flex():
                p.slide("a")
                p.slide("b")
            p.code("c", "x")
    page = _page(body)
    nodes = model.layout(page.blocks)
    row = nodes[0]
    assert [c.weight for c in row.children] == [0.5, 0.5]
    inner = row.children[0]
    assert [c.weight for c in inner.children] == [None, None]


def test_row_shares_renormalise_over_the_children_a_target_keeps():
    def body(p):
        p.title("Title")
        with p.flex("row", widths=[1, 2, 1]):
            p.slide("a")
            p.prose("book only")
            p.code("c", "x")
    page = _page(body)
    viewer = model.layout(model.select_blocks(page, "viewer"))[0]
    assert [c.kind for c in viewer.children] == ["slide", "code"]
    assert model.row_shares(viewer) == [0.5, 0.5]


def test_default_row_has_no_shares():
    page = _page(_row(None))
    row = model.layout(page.blocks)[0]
    assert model.row_shares(row) == [None, None]


@pytest.mark.parametrize("widths,match", [
    ([1], "1 entries, expected 2"),
    ([1, 2, 3], "3 entries, expected 2"),
    ([1, 0], "positive"),
    ([1, -2], "positive"),
    ([1, True], "positive"),
    ([1, float("inf")], "positive"),
])
def test_bad_widths_are_refused(widths, match):
    with pytest.raises(model.ValidationError, match=match):
        _page(_row(widths))


def test_widths_belong_to_a_row():
    p = PageBuilder()
    with pytest.raises(model.ValidationError, match="row"):
        p.flex("column", widths=[1, 2])


def test_a_child_with_its_own_width_cannot_take_a_share():
    def body(p):
        p.title("Title")
        with p.flex("row", widths=[1, 2]):
            p.slide("a")
            p.image("image.png", width_px=300)
    with pytest.raises(model.ValidationError, match="width of its own"):
        _page(body)


def test_a_weight_outside_a_row_is_refused():
    page = _page(_row([1, 3]))
    blocks = [b for b in page.blocks if b.kind not in (model.FLEX_KIND, model.END_KIND)]
    page = model.Page(id="p", title="T", blocks=blocks)
    with pytest.raises(model.ValidationError, match="not a child of a row"):
        model.check_layout(page)


def test_viewer_row_carries_each_childs_share():
    md = render_marp_page(_page(_row([1, 9])))
    assert md.count("--lk-grow:10.00") == 1
    assert md.count("--lk-grow:90.00") == 1
    assert "--lk-grow" not in render_marp_page(_page(_row(None)))


def test_pptx_row_splits_the_width_in_the_given_proportion(tmp_path):
    slide = _pptx(_row([1, 3]), tmp_path)
    one = next(s for s in slide.shapes if s.has_text_frame and s.text == "one")
    code = next(s for s in slide.shapes if s.has_text_frame and "two" in s.text)
    layout = Layout.from_ratio("16:9")
    gutter = px(24)
    assert code.left > one.left
    span = code.left - one.left
    available = layout.content_width - gutter
    assert abs(span - (available // 4 + gutter)) <= px(4)
