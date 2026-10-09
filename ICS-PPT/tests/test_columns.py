"""``p.columns``: slide text set side by side."""

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from lecturekit import Lecture, i18n, model
from lecturekit.model import Block, ValidationError, check_block
from lecturekit.renderers.latex.blocks import Ctx as LatexCtx, emit_block
from lecturekit.renderers.pptx import PptxRenderer
from lecturekit.renderers.transcript import build_html
from lecturekit.renderers.transcript.images import Embedder
from lecturekit.renderers.viewer import render_marp_page
from lecturekit.renderers.viewer.blocks import render_block


def _columns_block(items, widths=None):
    return Block(kind="columns", content={"items": items, "widths": widths})


def _lecture(body):
    lecture = Lecture(id="lec", title="L")
    lecture.page("p", body=body)
    return lecture.build()


def _block(body):
    return _lecture(body).children[0].blocks[0]


class ColumnsModelTest(unittest.TestCase):
    def test_columns_is_a_known_block_kind(self):
        self.assertIn("columns", model.BLOCK_KINDS)

    def test_check_block_accepts_two_columns(self):
        check_block(_columns_block(["a", "b"]), page_id="p")

    def test_check_block_rejects_a_single_column(self):
        with self.assertRaises(ValidationError):
            check_block(_columns_block(["a"]), page_id="p")

    def test_check_block_rejects_a_blank_column(self):
        with self.assertRaises(ValidationError):
            check_block(_columns_block(["a", "  "]), page_id="p")

    def test_check_block_rejects_widths_of_another_length(self):
        with self.assertRaises(ValidationError):
            check_block(_columns_block(["a", "b"], [1.0]), page_id="p")

    def test_a_mark_is_legal_in_a_column(self):
        check_block(_columns_block(["<mark>a</mark>", "b"]), page_id="p")


class ColumnsAuthoringTest(unittest.TestCase):
    def test_columns_are_stored_in_order(self):
        def body(p):
            p.title("W")
            p.columns("- left", "- right")

        block = _block(body)
        self.assertEqual(block.kind, "columns")
        self.assertEqual(block.content, {"items": ["- left", "- right"], "widths": None})

    def test_a_column_is_slide_text(self):
        def body(p):
            p.title("W")
            p.columns("headline\n- ==key== point", "other")

        self.assertEqual(
            _block(body).content["items"],
            ["**headline**\n- <mark>key</mark> point", "**other**"],
        )

    def test_autobold_false_leaves_prose_lines_alone(self):
        def body(p):
            p.title("W")
            p.columns("plain", "text", autobold=False)

        block = _block(body)
        self.assertEqual(block.content["items"], ["plain", "text"])
        self.assertFalse(block.autobold)

    def test_widths_are_normalized(self):
        def body(p):
            p.title("W")
            p.columns("a", "b", widths=[1, 3], autobold=False)

        self.assertEqual(_block(body).content["widths"], [0.25, 0.75])

    def test_widths_must_match_the_columns(self):
        def body(p):
            p.title("W")
            p.columns("a", "b", widths=[1, 2, 3])

        with self.assertRaises(ValidationError):
            _block(body)

    def test_one_column_raises_on_build(self):
        def body(p):
            p.title("W")
            p.columns("alone")

        with self.assertRaises(ValidationError):
            _block(body)

    def test_footnote_chains(self):
        def body(p):
            p.title("W")
            p.columns("a", "b").footnote("src")

        self.assertEqual(_block(body).footnotes, ("src",))


class ColumnsViewerTest(unittest.TestCase):
    def _lines(self, items, widths=None):
        return render_block(_columns_block(items, widths))

    def test_one_div_per_column_inside_the_track(self):
        html = "\n".join(self._lines(["- a", "- b", "- c"]))
        self.assertEqual(html.count('<div class="lk-columns">'), 1)
        self.assertEqual(html.count('<div class="lk-column"'), 3)

    def test_column_markdown_sits_between_blank_lines(self):
        lines = self._lines(["- a", "- b"])
        at = lines.index("- a")
        self.assertEqual((lines[at - 1], lines[at + 1]), ("", ""))

    def test_unweighted_columns_carry_no_inline_style(self):
        self.assertNotIn("style=", "\n".join(self._lines(["a", "b"])))

    def test_weighted_columns_carry_their_share(self):
        html = "\n".join(self._lines(["a", "b"], [0.25, 0.75]))
        self.assertIn('style="flex-grow:25.00"', html)
        self.assertIn('style="flex-grow:75.00"', html)

    def test_the_page_carries_the_columns(self):
        def body(p):
            p.title("W")
            p.columns("- left", "- right")

        page = _lecture(body).children[0]
        markdown = render_marp_page(page, slide_size=(1280, 720))
        self.assertIn('<div class="lk-columns">', markdown)
        self.assertIn("- right", markdown)


class ColumnsPptxTest(unittest.TestCase):
    def test_columns_are_text_boxes_side_by_side(self):
        def body(p):
            p.title("W")
            p.columns("- left one\n- left two", "- right", autobold=False)
            p.slide("below")

        with tempfile.TemporaryDirectory() as tmp:
            out = PptxRenderer().render(_lecture(body), Path(tmp))
            slide = Presentation(str(out)).slides[0]
            boxes = {s.text_frame.text.split("\n")[0]: s
                     for s in slide.shapes if s.has_text_frame}
        left = next(s for text, s in boxes.items() if "left one" in text)
        right = next(s for text, s in boxes.items() if "right" in text)
        below = next(s for text, s in boxes.items() if "below" in text)
        self.assertEqual(left.top, right.top)
        self.assertGreaterEqual(right.left, left.left + left.width)
        self.assertEqual(left.width, right.width)
        # The next block clears the taller column.
        self.assertGreaterEqual(below.top, left.top + left.height)


class ColumnsTranscriptTest(unittest.TestCase):
    def test_columns_print_in_order(self):
        def body(p):
            p.title("W")
            p.columns("left text", "right text", autobold=False)

        lecture = _lecture(body)
        html = build_html(lecture, Embedder(None, lecture.borrowed))
        self.assertLess(html.index("left text"), html.index("right text"))


class ColumnsLatexTest(unittest.TestCase):
    def test_columns_forced_into_the_book_print_in_order(self):
        ctx = LatexCtx(lecture_id="lec", page_id="p", slide_width=1280,
                       assets=None, asset_root=None, lang=None)
        tex = emit_block(_columns_block(["left text", "right text"]), ctx)
        self.assertLess(tex.index("left text"), tex.index("right text"))


class ColumnsI18nTest(unittest.TestCase):
    def setUp(self):
        def body(p):
            p.title("T")
            p.columns("左", "右")

        self.lecture = _lecture(body)

    def test_one_key_per_column(self):
        keys = [entry.key for entry in i18n.collect(self.lecture)]
        self.assertEqual(keys[-2:], ["p.columns.1.col.1", "p.columns.1.col.2"])

    def test_a_translation_follows_the_rules_of_slide_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = i18n.overlay_path(Path(tmp), "en")
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(
                f"[{i18n.toml_string('p.columns.1.col.2')}]\n"
                f"text = {i18n.toml_string('==right==')}\n",
                encoding="utf-8",
            )
            applied = i18n.apply(self.lecture, Path(tmp), "en")
        block = model.flatten_pages(applied.children)[0].blocks[0]
        self.assertEqual(
            block.content["items"], ["**左**", "**<mark>right</mark>**"],
        )


if __name__ == "__main__":
    unittest.main()
