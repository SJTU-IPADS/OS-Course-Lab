"""Uniform alignment, nested boxes, and their target-specific behaviour."""

from dataclasses import FrozenInstanceError

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from lecturekit import Lecture, model
from lecturekit.dsl import PageBuilder
from lecturekit.renderers.pptx import PptxRenderer
from lecturekit.renderers.pptx.layout import Layout, px
from lecturekit.renderers.viewer import render_marp_page
from lecturekit.serialize import lecture_to_dict


def _lecture(body):
    lecture = Lecture(id="layout", title="Layout")
    lecture.page("page", body=body)
    return lecture


def _page(body):
    return _lecture(body).build().children[0]


def _body(p):
    p.title("Title")
    p.slide("short", autobold=False)


def _pptx(body, tmp_path):
    Image.new("RGB", (200, 100), "white").save(tmp_path / "image.png")
    out = PptxRenderer(asset_root=tmp_path).render(_lecture(body).build(), tmp_path)
    return Presentation(str(out)).slides[0]


def _text(slide, text):
    return next(s for s in slide.shapes if s.has_text_frame and s.text == text)


def test_elements_have_immutable_shared_alignment():
    page = _page(_body)
    block = page.blocks[0]
    flex = model.Flex(children=(block,))
    for element in (page, block, flex):
        assert isinstance(element, model.LecturekitElement)
        moved = element.aligned(halign="right", valign="bottom")
        assert (moved.halign, moved.valign) == ("right", "bottom")
        assert (element.halign, element.valign) == (None, None)
        with pytest.raises(FrozenInstanceError):
            moved.halign = "left"
    with pytest.raises(model.ValidationError, match="axis"):
        block.aligned(align="left")


@pytest.mark.parametrize("axis,value", [("halign", "middle"), ("valign", "center")])
def test_bad_alignments_are_rejected_on_every_handle(axis, value):
    p = PageBuilder()
    handles = (p, p.slide("text"), p.flex())
    for handle in handles:
        with pytest.raises(model.ValidationError, match=axis):
            getattr(handle, axis)(value)


def test_nested_containers_preserve_leaf_order_and_handle_chaining():
    def body(p):
        p.title("Title")
        assert p.valign("middle").halign("center") is p
        with p.flex("row", valign="middle") as row:
            assert row.halign("left") is row
            p.slide("first").valign("bottom").footnote("source")
            with p.flex():
                p.code("c", "int x;").halign("right")
                p.slide("last")
    page = _page(body)
    root, = model.layout(page.blocks)
    assert root.direction == "row"
    assert root.halign == "left"
    assert root.children[0].valign == "bottom"
    assert root.children[0].footnotes == ("source",)
    assert root.children[1].children[0].halign == "right"
    assert root.children[1].children[1].content == "**last**"


@pytest.mark.parametrize("blocks,error", [
    ([model.Block("end", None)], "closes no"),
    ([model.Block("flex", {"direction": "column"})], "never closed"),
    ([model.Block("flex", {"direction": "column"}), model.Block("end", None)], "empty"),
    ([model.Block("flex", {"direction": "diagonal"}), model.Block("slide", "x"),
      model.Block("end", None)], "direction"),
])
def test_malformed_container_models_are_rejected(blocks, error):
    with pytest.raises(model.ValidationError, match=error):
        model.check_layout(model.Page(id="p", title="T", blocks=blocks))


@pytest.mark.parametrize("valign", ["middle", "bottom"])
def test_page_alignment_excludes_spreading_the_same_space(valign):
    def body(p):
        _body(p)
        p.valign(valign)
        p.gap()
    with pytest.raises(model.ValidationError, match="gap"):
        _page(body)


def test_side_columns_remain_page_level():
    def body(p):
        p.title("Title")
        with p.flex():
            p.side_image("image.png")
    with pytest.raises(model.ValidationError, match="cannot hold side_image"):
        _page(body)
    p = PageBuilder()
    with pytest.raises(model.ValidationError, match="page's alignment"):
        p.side_image("image.png").halign("left")


def test_empty_filtered_containers_leave_no_boxes_in_a_target():
    def body(p):
        p.title("Title")
        with p.flex():
            p.code("c", "hidden", only=["pptx"])
        p.slide("visible")
    md = render_marp_page(_page(body))
    assert "hidden" not in md
    assert "lk-flex" not in md
    assert "visible" in md


def test_viewer_boxes_keep_markdown_and_leaf_reveal_steps():
    def body(p):
        p.title("Title")
        with p.flex("row", valign="middle"):
            p.code("c", "int x;").halign("center").footnote("source")
            with p.flex():
                p.slide("- one\n- two", reveal="items")
                p.image("image.png", width_px=200)
    md = render_marp_page(_page(body), reveal=True)
    assert 'lk-flex-row' in md and 'lk-flex-column' in md
    assert 'align-items:center;justify-content:center' in md
    assert '\n\n```c\nint x;\n```\n' in md
    assert md.count('data-reveal="') == 3
    assert 'data-reveal-items="1"' in md
    assert 'source' in md
    assert md.count('<div ') == md.count('</div>')


def test_gap_seams_separate_top_level_elements_not_container_children():
    def body(p):
        p.title("Title")
        p.gap()
        with p.flex("row"):
            p.slide("one")
            p.slide("two")
        p.code("c", "three")
    md = render_marp_page(_page(body))
    assert md.count('class="lk-gap-spacer"') == 1
    assert md.count('class="lk-gap-block"') == 2


def test_alignment_survives_animation_and_serialization():
    def body(p):
        p.title("Title")
        p.valign("middle")
        with p.flex("row"):
            p.slide("fixed")
            p.frames("a.png", "b.png", width_px=200).halign("right")
        p.highlight("after").halign("left")
    lecture = _lecture(body).build()
    pages = model.flatten_pages(lecture.children)
    for page in pages:
        assert page.valign == "middle"
        assert next(b for b in page.blocks if b.kind == "image").halign == "right"
        assert len(model.layout(page.blocks)) == 2
    first = render_marp_page(pages[0], reveal=True)
    assert 'class="lk-box lk-held"' in first
    assert 'data-reveal="2"' not in first
    data = lecture_to_dict(lecture)
    assert data["children"][0]["valign"] == "middle"


@pytest.mark.parametrize("axis,value", [("halign", "center"), ("halign", "right"),
                                        ("valign", "middle"), ("valign", "bottom")])
def test_pptx_leaf_and_page_alignment_use_their_content_band(axis, value, tmp_path):
    def body(p):
        _body(p)
        if axis == "valign":
            p.valign(value)
        else:
            p.halign(value)
    slide = _pptx(body, tmp_path)
    text = _text(slide, "short")
    layout = Layout.from_ratio("16:9")
    if axis == "halign":
        slack = layout.content_width - text.width
        expected = layout.content_left + round(slack * model.ALIGN_FRACTION[value])
        assert abs(text.left - expected) <= 1
    else:
        plain = _pptx(_body, tmp_path)
        top = _text(plain, "short").top
        expected = top + round((layout.content_bottom - top - text.height)
                               * model.ALIGN_FRACTION[value])
        assert abs(text.top - expected) <= 1


def test_pptx_row_gives_fixed_image_width_and_centres_text_against_it(tmp_path):
    def body(p):
        p.title("Title")
        with p.flex("row", valign="middle"):
            p.slide("short", autobold=False)
            p.image("image.png", width_px=360)
    slide = _pptx(body, tmp_path)
    text = _text(slide, "short")
    pic = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert pic.width == px(360)
    assert abs(text.top + text.height / 2 - pic.top - pic.height / 2) <= 1
    assert text.left + text.width < pic.left


def test_pptx_nested_columns_share_row_height_and_child_override(tmp_path):
    def body(p):
        p.title("Title")
        with p.flex("row", valign="bottom"):
            with p.flex():
                p.slide("first", autobold=False)
                p.slide("second", autobold=False)
            p.code("c", "short").valign("top")
    slide = _pptx(body, tmp_path)
    first, second = _text(slide, "first"), _text(slide, "second")
    code = _text(slide, "short")
    assert first.top == code.top
    assert second.top > code.top
    assert first.left < code.left


def test_containers_are_transparent_on_paper(tmp_path):
    from lecturekit.renderers.transcript.blocks import render_block
    from lecturekit.renderers.transcript.images import Embedder

    def body(p):
        p.title("Title")
        with p.flex("row", halign="right"):
            p.slide("first")
            p.prose("book text")
            p.code("c", "last")
    page = _page(body)
    transcript = model.select_blocks(page, "transcript")
    assert [b.kind for b in transcript] == ["slide", "code"]
    latex_blocks = model.select_blocks(page, "latex")
    assert [b.kind for b in latex_blocks] == ["prose", "code"]
    printed = "".join(render_block(b, Embedder(tmp_path)) for b in transcript)
    assert printed.index("first") < printed.index("last")
    from lecturekit.renderers.latex.blocks import Ctx, emit_block
    from lecturekit.renderers.latex.assets import AssetCopier
    ctx = Ctx(lecture_id="layout", page_id="page", slide_width=1280,
              assets=AssetCopier(tmp_path))
    book = "".join(emit_block(b, ctx) for b in latex_blocks)
    assert book.index("book text") < book.index("last")


def test_pptx_alignment_moves_figure_and_caption_together(tmp_path):
    def body(p):
        p.title("Title")
        p.image("image.png", width_px=200, caption="Caption").halign("right")
    slide = _pptx(body, tmp_path)
    pic = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    caption = _text(slide, "Caption")
    layout = Layout.from_ratio("16:9")
    assert pic.left + pic.width == layout.content_left + layout.content_width
    assert abs(pic.left + pic.width / 2 - caption.left - caption.width / 2) <= 1
    assert caption.top > pic.top + pic.height


def test_pptx_aligned_animation_reserves_the_finished_layout(tmp_path):
    def body(p):
        p.title("Title")
        p.valign("middle")
        with p.flex("row", valign="middle"):
            p.slide("fixed")
            p.frames("image.png", "image.png", width_px=200)
        p.highlight("later")
    Image.new("RGB", (200, 100), "white").save(tmp_path / "image.png")
    out = PptxRenderer(asset_root=tmp_path).render(_lecture(body).build(), tmp_path)
    slides = Presentation(str(out)).slides
    pics = [next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
            for slide in slides]
    assert (pics[0].left, pics[0].top, pics[0].width, pics[0].height) == (
        pics[1].left, pics[1].top, pics[1].width, pics[1].height)
    assert all(not (s.has_text_frame and s.text == "later") for s in slides[0].shapes)
    assert any(s.has_text_frame and s.text == "later" for s in slides[1].shapes)


def test_pptx_column_alignment_does_not_move_a_child_override(tmp_path):
    def body(p):
        p.title("Title")
        with p.flex(halign="center"):
            p.slide("longer centered text", autobold=False)
            p.slide("left", autobold=False).halign("left")
    slide = _pptx(body, tmp_path)
    layout = Layout.from_ratio("16:9")
    assert _text(slide, "left").left == layout.content_left
    center = _text(slide, "longer centered text")
    assert abs(center.left + center.width / 2 - layout.width / 2) <= 1


def test_viewer_layout_preserves_side_metadata_and_slide_coordinate_annotations():
    def body(p):
        p.title("Title")
        p.halign("left")
        p.side_image("image.png").footnote("image source").annotate("side note")
        p.slide("text").annotate("text note")
    md = render_marp_page(_page(body), reveal=True)
    assert 'image source' in md
    assert 'side note' in md and 'text note' in md
    assert 'data-reveal="0" class="annotation' in md
    assert md.index('side note') > md.index('</div>')
    assert md.count('class="lk-box') == 1
