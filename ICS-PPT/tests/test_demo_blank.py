"""A demo block with a blank command: file buttons, and nothing to run.

`p.demo("查看源码", "", files=[...])` puts a source file one press away on a
page that explains a listing. There is no command, so no target draws a run
button or a `$` line, and the deck's table of runnable commands leaves the
block out.
"""

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from lecturekit import Lecture, demo
from lecturekit.renderers.latex.blocks import Ctx, emit_block
from lecturekit.renderers.pptx import PptxRenderer
from lecturekit.renderers.viewer.blocks import render_block


def demo_lecture(command="", name="show the source", **kw):
    lecture = Lecture(id="lec", title="L")

    def body(p):
        p.title("P")
        p.demo(name, command, files=["examples/main.c"], **kw)

    lecture.page("p1", body=body)
    return lecture


def demo_block(command="", **kw):
    blocks = demo_lecture(command, **kw).build().children[0].blocks
    return next(b for b in blocks if b.kind == "demo")


def slide_texts(lecture):
    out = PptxRenderer().render(lecture.build(), Path(tempfile.mkdtemp()))
    return [
        para.text
        for shape in Presentation(str(out)).slides[0].shapes
        if shape.has_text_frame
        for para in shape.text_frame.paragraphs
    ]


class BlankCommandTest(unittest.TestCase):
    def test_a_blank_command_has_no_prompt_lines(self):
        for blank in ("", "   ", "\n", " \n\t"):
            with self.subTest(command=blank):
                self.assertEqual(demo.prompt_lines(blank), [])

    def test_a_blank_command_is_not_runnable(self):
        for blank in ("", "   "):
            with self.subTest(command=blank):
                self.assertEqual(demo.collect(demo_lecture(blank).build()), {})

    def test_a_command_beside_it_is_still_runnable(self):
        lecture = demo_lecture("")

        def body(p):
            p.title("Q")
            p.demo("build", "make")

        lecture.page("p2", body=body)
        table = demo.collect(lecture.build())
        self.assertEqual([spec.command for spec in table.values()], ["make"])


class BlankCommandDeckTest(unittest.TestCase):
    def test_the_deck_draws_the_file_button_and_the_name(self):
        html = render_block(demo_block())[0]
        self.assertIn('class="lk-demo-file" type="button" disabled', html)
        self.assertIn(">main.c</button>", html)
        self.assertIn('<span class="lk-demo-name">show the source</span>', html)

    def test_the_deck_draws_no_run_button_and_no_command_line(self):
        for blank in ("", "   "):
            with self.subTest(command=blank):
                html = render_block(demo_block(blank))[0]
                self.assertNotIn("lk-demo-run", html)
                self.assertNotIn("lk-demo-cmd", html)
                self.assertNotIn("$", html)

    def test_the_deck_keeps_it_a_one_line_chip(self):
        html = render_block(demo_block())[0]
        self.assertIn('data-lk-demo-form="inline"', html)
        self.assertNotIn("lk-demo-term", html)

    def test_the_deck_still_draws_the_description(self):
        html = render_block(demo_block(description="the caller"))[0]
        self.assertIn('<span class="lk-demo-desc">the caller</span>', html)

    def test_the_deck_still_draws_a_recorded_output(self):
        html = render_block(demo_block(output="42"))[0]
        self.assertIn('data-lk-demo-form="block"', html)
        self.assertIn('<pre class="lk-demo-term"><code class="lk-demo-out">42</code>', html)
        self.assertNotIn("lk-demo-run", html)
        self.assertNotIn("lk-demo-cmd", html)

    def test_a_command_keeps_its_run_button_beside_the_file_button(self):
        html = render_block(demo_block("make"))[0]
        self.assertLess(html.index("lk-demo-run"), html.index("lk-demo-file"))
        self.assertIn('<code class="lk-demo-cmd">$ make</code>', html)


class BlankCommandStillTargetsTest(unittest.TestCase):
    def test_the_pptx_prints_the_name_and_no_listing(self):
        # The title and the caption; the listing box and its accent bar are not drawn.
        texts = slide_texts(demo_lecture(description="the caller"))
        self.assertEqual(texts, ["P", "show the source — the caller"])
        with_command = slide_texts(demo_lecture("make", description="the caller"))
        self.assertEqual(with_command[:3], texts + ["$ make"])

    def test_the_pptx_draws_nothing_for_a_blank_command_without_a_name(self):
        self.assertEqual(slide_texts(demo_lecture(name="")), ["P"])

    def test_the_book_prints_the_box_without_a_listing(self):
        ctx = Ctx(lecture_id="lec", page_id="p1", slide_width=1280, assets=None)
        out = emit_block(demo_block(description="the caller"), ctx)
        self.assertIn(r"\begin{tcolorbox}", out)
        self.assertIn(r"\textbf{show the source}\\", out)
        self.assertIn("the caller", out)
        self.assertNotIn("lstlisting", out)
        self.assertTrue(out.rstrip().endswith(r"\end{tcolorbox}"))

    def test_the_book_prints_a_recorded_output_without_a_prompt(self):
        ctx = Ctx(lecture_id="lec", page_id="p1", slide_width=1280, assets=None)
        out = emit_block(demo_block("  ", output="42"), ctx)
        self.assertIn("\\begin{lstlisting}\n42\n\\end{lstlisting}", out)
        self.assertNotIn("$", out)


if __name__ == "__main__":
    unittest.main()
