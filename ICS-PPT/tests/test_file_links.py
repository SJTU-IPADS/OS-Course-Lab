"""A file button in a line of text: `[label](file:path)`.

`p.demo(..., files=[...])` hangs a file on a command. A sentence that only
mentions a file names it where it is mentioned, and the properties worth
protecting are the ones the demo's button has: the page carries an id and never
a path, the table of ids is rebuilt from what the lecture still says, and a
target that cannot open a panel prints the label as the word it is.
"""

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from lecturekit import filelinks, model, source
from lecturekit.dsl import Lecture
from lecturekit.renderers.latex.text import inline as latex_inline
from lecturekit.renderers.pptx import PptxRenderer
from lecturekit.renderers.pptx.text import parse_inline
from lecturekit.renderers.transcript.text import inline as sheet_inline
from lecturekit.renderers.viewer.blocks import render_block

ASSETS = Path(__file__).parent.parent / "lecturekit" / "renderers" / "viewer" / "assets"
THEME = Path(__file__).parent.parent / "themes" / "basic-office.css"


def lecture_of(body) -> Lecture:
    lecture = Lecture(id="lec", title="L")

    def page(p):
        p.title("P")
        body(p)

    lecture.page("p1", body=page)
    return lecture


def blocks_of(body, kind: str) -> list[model.Block]:
    built = lecture_of(body).build()
    return [b for b in built.children[0].blocks if b.kind == kind]


def html_of(body, kind: str = "slide") -> str:
    """What the deck's markdown holds for the page's blocks of one kind."""
    return "\n".join(
        line for block in blocks_of(body, kind) for line in render_block(block)
    ).strip()


class FileLinkSyntaxTest(unittest.TestCase):
    """What counts as a file link, and where one is looked for."""

    def test_a_link_gives_its_label_and_its_path(self):
        self.assertEqual(
            filelinks.find("see [main.c](file:examples/main.c) and [h](file:a.h)"),
            [("main.c", "examples/main.c"), ("h", "a.h")],
        )

    def test_an_ordinary_link_is_not_a_file_link(self):
        self.assertEqual(filelinks.find("[site](https://example.com/file:x)"), [])
        self.assertIsNone(filelinks.path_of("https://example.com"))
        self.assertEqual(filelinks.path_of("file:a/b.c"), "a/b.c")

    def test_a_code_span_or_a_fence_shows_the_syntax_and_is_not_read(self):
        text = "write `[a.c](file:a.c)`\n```\n[b.c](file:b.c)\n```\n[c.c](file:c.c)"
        self.assertEqual(filelinks.find(text), [("c.c", "c.c")])
        self.assertEqual(filelinks.errors(text, allowed=False),
                         filelinks.errors("[c.c](file:c.c)", allowed=False))

    def test_a_label_may_be_a_code_span(self):
        self.assertEqual(filelinks.find("[`a.c`](file:a.c)"), [("`a.c`", "a.c")])

    def test_replace_rewrites_the_links_and_nothing_else(self):
        out = filelinks.replace(
            "x [a](file:p/a.c) y `[b](file:b.c)` z", lambda label, path: f"<{label}|{path}>"
        )
        self.assertEqual(out, "x <a|p/a.c> y `[b](file:b.c)` z")

    def test_text_without_a_link_comes_back_as_it_was(self):
        text = "nothing here: profile:x, [a](b)"
        self.assertEqual(filelinks.replace(text, lambda *_: "!"), text)
        self.assertEqual(filelinks.errors(text, allowed=False), [])

    def test_a_path_must_stay_inside_the_lecture(self):
        for path in ("../other/a.c", "a/../../b.c", "///abs", "//abs", "/",
                     "/../a.c", "/lectures/../../a.c"):
            with self.subTest(path=path):
                problems = filelinks.errors(f"[x](file:{path})", allowed=True)
                self.assertEqual(len(problems), 1)
                self.assertIn("relative to the lecture directory", problems[0])

    def test_a_near_miss_is_an_error_and_not_text_on_the_slide(self):
        for text in ("[](file:a.c)", "[a.c](file: a.c)", "[a.c](File:a.c)",
                     "[a.c](file:)", "[a.c]( file:a.c)"):
            with self.subTest(text=text):
                problems = filelinks.errors(text, allowed=True)
                self.assertEqual(len(problems), 1)
                self.assertIn("malformed file link", problems[0])

    def test_a_leading_slash_names_a_file_of_the_project_root(self):
        text = "[seq.c](file:/lectures/mini-cpu/src/lab2/seq.c)"
        self.assertEqual(
            filelinks.find(text), [("seq.c", "/lectures/mini-cpu/src/lab2/seq.c")]
        )
        self.assertEqual(filelinks.errors(text, allowed=True), [])
        self.assertTrue(filelinks.is_rooted("/lectures/a.c"))
        self.assertFalse(filelinks.is_rooted("examples/a.c"))

    def test_files_and_links_share_one_path_rule(self):
        for path in ("//abs/a.c", "../a.c", "/../a.c"):
            with self.subTest(path=path):
                self.assertIsNotNone(filelinks.path_problem(path))
                with self.assertRaises(model.ValidationError):
                    source.normalize([path])
        for path in ("examples/a.c", "/lectures/lab/a.c"):
            with self.subTest(path=path):
                self.assertIsNone(filelinks.path_problem(path))
                self.assertEqual(source.normalize([path]), [path])


class FileLinkPlacementTest(unittest.TestCase):
    """Slide text and table cells may carry one; nothing else may."""

    def test_slide_columns_and_table_accept_a_link(self):
        def body(p):
            p.slide("see [a.c](file:a.c)")
            p.columns("left [b.c](file:b.c)", "right")
            p.table([["[c.c](file:c.c)", "x"]], headers=["file", "what"])

        lecture_of(body).build()

    def test_other_text_refuses_a_link(self):
        link = "[a.c](file:a.c)"
        cases = {
            "notes": lambda p: p.notes(f"see {link}"),
            "prose": lambda p: p.prose(f"see {link}"),
            "aside": lambda p: p.aside(f"see {link}"),
            "highlight": lambda p: p.highlight(f"see {link}"),
            "demo description": lambda p: p.demo("run", "make", description=link),
            "footnote": lambda p: p.slide("text").footnote(f"see {link}"),
        }
        for name, body in cases.items():
            with self.subTest(where=name):
                with self.assertRaisesRegex(model.ValidationError, "is allowed in"):
                    lecture_of(body).build()

    def test_code_and_a_recorded_output_are_verbatim(self):
        def body(p):
            p.code("markdown", "[a.c](file:/abs/a.c)")
            p.demo("run", "cat notes.md", output="[a.c](file:/abs/a.c)")

        lecture_of(body).build()

    def test_a_bad_path_names_the_page(self):
        with self.assertRaisesRegex(model.ValidationError, r"'\.\./a\.c'.* on p1"):
            lecture_of(lambda p: p.slide("see [a.c](file:../a.c)")).build()

    def test_a_malformed_link_is_refused(self):
        with self.assertRaisesRegex(model.ValidationError, "malformed file link"):
            lecture_of(lambda p: p.slide("see [a.c](file: a.c)")).build()


class FileLinkTableTest(unittest.TestCase):
    """`sources.json` lists what the text names, as it lists a demo's files."""

    def test_links_join_the_table_beside_a_demos_files(self):
        def body(p):
            p.slide("see [a.c](file:examples/a.c)")
            p.columns("left", "right [b.h](file:examples/b.h)")
            p.table([["[Makefile](file:Makefile)"]], headers=["file"])
            p.demo("build", "make", files=["examples/a.c", "examples/c.c"])

        table = source.collect(lecture_of(body).build())
        self.assertEqual(
            sorted(entry.path for entry in table.values()),
            ["Makefile", "examples/a.c", "examples/b.h", "examples/c.c"],
        )
        self.assertEqual(table[source.source_id("examples/b.h")].path, "examples/b.h")

    def test_a_disabled_block_names_nothing(self):
        def body(p):
            p.slide("see [a.c](file:a.c)").disable()

        self.assertEqual(source.collect(lecture_of(body).build()), {})

    def test_a_link_shown_as_code_names_nothing(self):
        def body(p):
            p.slide("write `[a.c](file:a.c)` in the text")

        self.assertEqual(source.collect(lecture_of(body).build()), {})


class FileLinkButtonTest(unittest.TestCase):
    """The deck draws the label as a button carrying the file's id."""

    def test_the_button_carries_the_id_and_the_label(self):
        html = html_of(lambda p: p.slide("see [the source](file:examples/main.c) here"))
        self.assertIn(f'data-lk-source="{source.source_id("examples/main.c")}"', html)
        self.assertIn(">the source</button>", html)
        self.assertNotIn("file:", html)
        self.assertNotIn("](", html)

    def test_a_file_of_the_project_root_is_a_button_too(self):
        path = "/lectures/mini-cpu/src/lab2/seq.c"
        html = html_of(lambda p: p.slide(f"see [seq.c](file:{path}) here"))
        self.assertIn(f'data-lk-source="{source.source_id(path)}"', html)
        self.assertIn(">seq.c</button>", html)

    def test_the_button_ships_disabled(self):
        html = html_of(lambda p: p.slide("see [a.c](file:a.c)"))
        self.assertIn('<button class="lk-text-file" type="button" disabled', html)

    def test_the_button_stays_in_its_sentence(self):
        html = html_of(lambda p: p.slide("- before [a.c](file:a.c) after", autobold=False))
        self.assertEqual(html.count("\n"), 0)
        self.assertTrue(html.startswith("- before <button"))
        self.assertTrue(html.endswith("</button> after"))

    def test_a_headline_is_still_bold_around_the_button(self):
        html = html_of(lambda p: p.slide("see [a.c](file:a.c)"))
        self.assertTrue(html.startswith("**see <button"))
        self.assertTrue(html.endswith("</button>**"))

    def test_markdown_does_not_read_the_label_a_second_time(self):
        # `__init__` between two inline tags is strong emphasis to markdown-it.
        html = html_of(lambda p: p.slide("- [__init__.py](file:pkg/__init__.py)"))
        self.assertNotIn("__", html)
        self.assertIn("&#95;&#95;init&#95;&#95;.py</button>", html)

    def test_a_code_span_label_is_code_in_the_button(self):
        html = html_of(lambda p: p.slide("- [`a.c`](file:a.c)"))
        self.assertIn("<code>a.c</code></button>", html)

    def test_a_label_cannot_smuggle_markup(self):
        html = html_of(lambda p: p.slide("- [<b onclick=x>a](file:a.c)"))
        self.assertIn("&lt;b onclick=x&gt;a</button>", html)

    def test_columns_and_table_cells_draw_the_button_too(self):
        columns = html_of(lambda p: p.columns("l [a.c](file:a.c)", "r"), "columns")
        self.assertIn(f'data-lk-source="{source.source_id("a.c")}"', columns)
        table = html_of(
            lambda p: p.table([["[b.c](file:b.c)", "x"]], headers=["file", "what"]),
            "table",
        )
        self.assertIn(f'data-lk-source="{source.source_id("b.c")}"', table)
        self.assertIn("</button> | x |", table)

    def test_a_link_shown_as_code_is_left_alone(self):
        html = html_of(lambda p: p.slide("- write `[a.c](file:a.c)`"))
        self.assertNotIn("<button", html)
        self.assertIn("`[a.c](file:a.c)`", html)

    def test_an_ordinary_link_is_left_to_markdown(self):
        html = html_of(lambda p: p.slide("- [site](https://example.com)"))
        self.assertEqual(html, "- [site](https://example.com)")

    def test_the_button_and_the_demo_button_agree_on_the_id(self):
        def body(p):
            p.slide("see [a.c](file:examples/a.c)")
            p.demo("build", "make", files=["examples/a.c"])

        wanted = f'data-lk-source="{source.source_id("examples/a.c")}"'
        self.assertIn(wanted, html_of(body, "slide"))
        self.assertIn(wanted, html_of(body, "demo"))


class FileLinkControllerTest(unittest.TestCase):
    """`source.js` arms the button, and the theme draws it."""

    def test_the_controller_arms_every_file_button(self):
        script = (ASSETS / "source.js").read_text(encoding="utf-8")
        self.assertIn('querySelectorAll("button[data-lk-source]")', script)

    def test_a_press_on_a_button_in_the_text_is_not_a_click_elsewhere(self):
        # With the panel up, a click outside the deck's chrome puts it away
        # before anything else hears it; the button must be let through.
        script = (ASSETS / "source.js").read_text(encoding="utf-8")
        self.assertIn('closest("[data-lk-chrome], .lk-demo, [data-lk-source]")', script)

    def test_the_theme_styles_the_button(self):
        css = THEME.read_text(encoding="utf-8")
        self.assertIn(".lk-text-file {", css)
        self.assertIn(".lk-text-file[disabled] {", css)


class FileLinkElsewhereTest(unittest.TestCase):
    """A target with no panel to open prints the label as text."""

    def test_pptx_runs_are_the_label_without_a_hyperlink(self):
        runs = parse_inline("see [main.c](file:examples/main.c) here")
        self.assertEqual("".join(run.text for run in runs), "see main.c here")
        self.assertTrue(all(run.link is None for run in runs))

    def test_pptx_keeps_the_emphasis_around_and_inside_the_label(self):
        runs = parse_inline("**see [`main.c`](file:examples/main.c)**")
        self.assertEqual("".join(run.text for run in runs), "see main.c")
        self.assertTrue(all(run.bold for run in runs))
        self.assertTrue(runs[-1].code)

    def test_pptx_still_links_an_ordinary_address(self):
        runs = parse_inline("[site](https://example.com)")
        self.assertEqual(runs[0].link, "https://example.com")

    def test_the_pptx_slide_prints_the_label(self):
        def body(p):
            p.slide("see [main.c](file:examples/main.c) here", autobold=False)
            p.table([["[Makefile](file:Makefile)", "rules"]], headers=["file", "what"])

        out = PptxRenderer().render(lecture_of(body).build(), Path(tempfile.mkdtemp()))
        texts: list[str] = []
        for shape in Presentation(str(out)).slides[0].shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                texts.extend(
                    cell.text for row in shape.table.rows for cell in row.cells
                )
        joined = "\n".join(texts)
        self.assertIn("see main.c here", joined)
        self.assertIn("Makefile", joined)
        self.assertNotIn("file:", joined)

    def test_latex_prints_the_label_without_an_href(self):
        out = latex_inline("see [main_light.c](file:src/main_light.c) here")
        self.assertEqual(out, r"see main\_light.c here")
        self.assertIn(r"\href", latex_inline("[site](https://example.com)"))

    def test_latex_sets_a_code_span_label_in_typewriter(self):
        self.assertEqual(latex_inline("[`a.c`](file:a.c)"), r"\texttt{a.c}")

    def test_the_sheet_prints_the_label_without_an_anchor(self):
        self.assertEqual(
            sheet_inline("see [main.c](file:examples/main.c) here"), "see main.c here"
        )
        self.assertIn("<a href=", sheet_inline("[site](https://example.com)"))


if __name__ == "__main__":
    unittest.main()
