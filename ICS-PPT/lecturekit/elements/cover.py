"""The `cover` element."""
from __future__ import annotations

from ..model import Block, element
from .. import model


@element
class Cover(Block):
    """The cover page: title, author, time, logo."""
    KIND = "cover"
    TARGETS = frozenset({"viewer", "pptx"})
    IN_CONTAINER = False
    REVEAL_STEP = False

    def check_content(self, page_id: str) -> None:
        content = self.content
        for field_name in ("author", "time"):
            value = content.get(field_name)
            if value is not None and (not isinstance(value, str) or not value.strip()):
                raise model.ValidationError(
                    f"Empty cover {field_name} on {page_id}"
                )
        logo = content.get("logo")
        if logo is None:
            return
        if isinstance(logo, str):
            if not logo.strip():
                raise model.ValidationError(f"Empty cover logo on {page_id}")
            return
        if not isinstance(logo, dict):
            raise model.ValidationError(
                f"Cover logo on {page_id} must be a string or left/right mapping"
            )
        for key in logo:
            if key not in ("left", "right"):
                raise model.ValidationError(
                    f"Cover logo on {page_id} has unknown slot {key!r}"
                )
        for key in ("left", "right"):
            value = logo.get(key)
            if value is not None and (not isinstance(value, str) or not value.strip()):
                raise model.ValidationError(
                    f"Empty cover logo {key} on {page_id}"
                )

    def to_marp(self, ctx):
        return []

    def to_pptx(self, ctx):
        from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
        from ..renderers.pptx import theme
        from ..renderers.pptx.blocks import (
            COVER_ACCENT,
            COVER_META_PT,
            COVER_TIME_PT,
            COVER_TITLE_PT,
            draw_cover_logos,
            set_line_spacing,
            style_run,
        )
        cover = self.content
        draw_cover_logos(cover.get("logo"), ctx)

        title_width = int(ctx.layout.width * 0.78)
        title_left = (ctx.layout.width - title_width) // 2
        title_top = int(ctx.layout.height * 0.24)
        title_height = int(ctx.layout.height * 0.26)
        box = ctx.slide.shapes.add_textbox(title_left, title_top, title_width, title_height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_line_spacing(p, COVER_TITLE_PT)
        style_run(p.add_run(), ctx.title, size_pt=COVER_TITLE_PT,
                   color=COVER_ACCENT)

        meta_top = int(ctx.layout.height * 0.57)
        meta_height = int(ctx.layout.height * 0.2)
        meta = ctx.slide.shapes.add_textbox(
            title_left, meta_top, title_width, meta_height
        )
        mtf = meta.text_frame
        mtf.word_wrap = True
        mtf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
        mtf.margin_left = mtf.margin_right = mtf.margin_top = mtf.margin_bottom = 0
        first = True
        if cover.get("author"):
            p = mtf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            set_line_spacing(p, COVER_META_PT)
            style_run(p.add_run(), cover["author"], size_pt=COVER_META_PT,
                       color=theme.FG, bold=True)
            first = False
        if cover.get("time"):
            p = mtf.paragraphs[0] if first else mtf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            set_line_spacing(p, COVER_TIME_PT)
            style_run(p.add_run(), cover["time"], size_pt=COVER_TIME_PT,
                       color=theme.DK2)
