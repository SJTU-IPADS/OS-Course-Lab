"""The `highlight` element."""
from __future__ import annotations

from ..model import Block, element
from .. import model


@element
class Highlight(Block):
    """A chip that sets one statement apart."""
    KIND = "highlight"
    TARGETS = frozenset({"viewer", "pptx", "latex", "transcript"})
    CENTERED = True

    def check_content(self, page_id: str) -> None:
        content = self.content
        text = content.get("text")
        if not isinstance(text, str) or not text.strip():
            raise model.ValidationError(f"Empty highlight text on {page_id}")
        tone = content.get("tone")
        if tone not in model.HIGHLIGHT_TONES:
            raise model.ValidationError(
                f"Highlight tone on {page_id} must be one of "
                f"{sorted(model.HIGHLIGHT_TONES)}, got {tone!r}"
            )

    def to_marp(self, ctx):
        """A centered chip: the outer element centers, the inner <span> is the wash.

        The split is the point — an inline-block span hugs the widest line instead
        of becoming a full-width banner. `inline_md` already turns newlines into
        <br>, so a multi-line highlight grows the same chip taller.

        A chip carrying math swaps the outer <p> for a blocked-out <span>, because
        Marp renders math while parsing markdown and a line starting with `<p` is a
        raw HTML block it never parses — a `<span>` with text after it on the line
        is not, so the row arrives as a paragraph and the formula is typeset. The
        box is the same one; only the tag it is drawn on differs.
        """
        from .. import model
        from ..renderers.viewer.blocks import CHIP_MATH_RE, chip_math_md, inline_md
        content = self.content
        lines = model.highlight_lines(content["text"])
        tone = content["tone"]
        classes = f"lk-highlight lk-highlight--{tone}"
        if any(CHIP_MATH_RE.search(line) for line in lines):
            body = "<br>".join(chip_math_md(line) for line in lines)
            return [f'<span class="{classes}"><span>{body}</span></span>', ""]
        body = inline_md("\n".join(lines))
        return [f'<p class="{classes}"><span>{body}</span></p>', ""]

    def footnote_marker(self, lines, numbers):
        from ..renderers.viewer import blocks
        return blocks.mark_highlight(lines, numbers)

    def to_pptx(self, ctx):
        """A centered chip: a hairline frame in the tone, holding text in that tone.

        PowerPoint has no inline-block, so the width is *estimated* from the longest
        line (like every other box in this renderer) and centered in the content
        column. As in the deck the shape carries no fill — the tone is the outline
        and the ink, and what is inside the frame is white slide.
        """
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Emu, Pt
        from .. import model
        from ..renderers.pptx import theme
        from ..renderers.pptx.blocks import (
            GAP,
            HIGHLIGHT_PAD,
            HIGHLIGHT_RULE_PT,
            HIGHLIGHT_SHADOW,
            _line,
            _shape,
            apply_runs,
            drop_shadow,
            set_line_spacing,
        )
        from ..renderers.pptx.layout import estimate_text_width, LINE_HEIGHT
        from ..renderers.pptx.text import parse_markdown, Run
        content = self.content
        lines = model.highlight_lines(content["text"])
        rows = [parse_markdown(line)[0].runs if line.strip() else [Run("")]
                for line in lines]
        size = theme.HIGHLIGHT_PT
        # Measured on the parsed runs, not the source: a `$$…$$` row is typeset as
        # an equation, so what the chip has to be wide enough for is the formula's
        # plain-text rendering, not the LaTeX that spelled it.
        width = min(
            ctx.layout.content_width,
            max(estimate_text_width("".join(run.text for run in runs), font_pt=size)
                for runs in rows)
            + 2 * HIGHLIGHT_PAD,
        )
        height = Emu(Pt(len(lines) * size * LINE_HEIGHT)) + 2 * HIGHLIGHT_PAD
        left = ctx.layout.content_left + (ctx.layout.content_width - width) // 2

        top = ctx.cursor.place(height)
        shape = _shape(ctx, MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.background()
        _line(shape, theme.CHIP_INK[content["tone"]], HIGHLIGHT_RULE_PT)
        drop_shadow(shape, theme.SIDENOTE_BORDER, **HIGHLIGHT_SHADOW)

        tf = shape.text_frame
        tf.word_wrap = False
        for i, runs in enumerate(rows):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = PP_ALIGN.CENTER
            set_line_spacing(para, size)
            apply_runs(para, runs, size_pt=size,
                        color=theme.CHIP_INK[content["tone"]])
            for run in para.runs:
                run.font.bold = True
        ctx.cursor.advance(GAP)

    def to_latex(self, ctx):
        """A centered chip. Lines become tabular rows, so it grows taller, not wider.

        A plain ``\\footnote`` after the macro would land outside the ``center``
        group and open a paragraph of its own, dropping the mark in the margin. So
        the marks ride the macro's optional argument — just outside the chip, as in
        the deck — and the texts follow, the same split ``figure`` uses.
        """
        from .. import model
        from ..renderers.latex.blocks import HIGHLIGHT_COLORS, display_math_inline
        from ..renderers.latex.text import inline
        content = self.content
        color = HIGHLIGHT_COLORS[content["tone"]]
        body = r" \\ ".join(
            inline(display_math_inline(line), refs=ctx.resolve_ref)
            for line in model.highlight_lines(content["text"])
        )
        marks = r"\footnotemark" * len(self.footnotes)
        out = r"\lkhighlight%s{%s}{%s}" % (f"[{marks}]" if marks else "", color, body)
        for note in self.footnotes:
            out += "\n" + r"\footnotetext{%s}" % inline(note, refs=ctx.resolve_ref)
        return out

    def to_transcript(self, embedder):
        from .. import model
        from ..renderers.transcript.text import inline
        tone = self.content.get("tone", "yellow")
        lines = model.highlight_lines(self.content["text"])
        body = "<br>".join(inline(line) for line in lines if line)
        return f'<p class="tx-highlight tx-{tone}"><span>{body}</span></p>'
