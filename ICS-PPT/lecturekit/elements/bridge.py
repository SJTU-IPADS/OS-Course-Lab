"""The `bridge` element."""
from __future__ import annotations

from ..model import Block, element
from .. import model


@element
class Bridge(Block):
    """A transition page: one to three short lines."""
    KIND = "bridge"
    TARGETS = frozenset({"viewer", "pptx"})
    IN_CONTAINER = False
    REVEAL_STEP = False

    def check_content(self, page_id: str) -> None:
        """Bridge text is plain, short, and carries nothing else.

        No footnotes or annotations: a bridge with something to source or point at
        has outgrown being a bridge and should be a normal page.
        """
        content = self.content if isinstance(self.content, dict) else {}
        text = content.get("text")
        if not isinstance(text, str) or not text.strip():
            raise model.ValidationError(f"Empty bridge text on {page_id}")
        lines = [line for line in (l.strip() for l in text.splitlines()) if line]
        if len(lines) > model.BRIDGE_MAX_LINES:
            raise model.ValidationError(
                f"a bridge holds at most {model.BRIDGE_MAX_LINES} lines on {page_id}, "
                f"got {len(lines)} — a longer transition wants a normal page"
            )
        if self.footnotes or self.annotations:
            raise model.ValidationError(
                f"a bridge carries no footnotes or annotations: {page_id}"
            )

    def to_marp(self, ctx):
        """A transition page's centered lines. Plain text — escaped, no inline md."""
        from ..renderers.viewer.blocks import escape_html
        body = "<br>".join(
            escape_html(line) for line in str(self.content["text"]).splitlines()
        )
        return [f'<div class="lk-bridge-body">{body}</div>', ""]

    def to_pptx(self, ctx):
        """A transition page: its plain-text lines centered on an empty slide.

        Like ``draw_cover``, this owns the whole slide — no title, no footnotes —
        so the renderer calls it instead of the block loop.
        """
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from ..renderers.pptx import theme
        from ..renderers.pptx.blocks import BRIDGE_PT, set_line_spacing, style_run
        width = int(ctx.layout.width * 0.82)
        left = (ctx.layout.width - width) // 2
        box = ctx.slide.shapes.add_textbox(left, 0, width, ctx.layout.height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        first = True
        for line in str(self.content["text"]).splitlines():
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = PP_ALIGN.CENTER
            set_line_spacing(p, BRIDGE_PT)
            style_run(p.add_run(), line, size_pt=BRIDGE_PT, color=theme.FG)
