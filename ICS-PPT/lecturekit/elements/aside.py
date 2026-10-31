"""The `aside` element."""
from __future__ import annotations

from ..model import Block, element


@element
class Aside(Block):
    """A quoted aside."""
    KIND = "aside"
    TARGETS = frozenset({"viewer", "pptx", "latex", "transcript"})

    def to_marp(self, ctx):
        return [f"> {str(self.content).strip()}", ""]

    def to_pptx(self, ctx):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Pt
        from ..renderers.pptx import theme
        from ..renderers.pptx.blocks import (
            GAP,
            _shape,
            apply_runs,
            no_line,
            set_line_spacing,
            solid,
            textbox,
        )
        from ..renderers.pptx.layout import estimate_text_height
        from ..renderers.pptx.text import parse_markdown, Run
        text = str(self.content).strip()
        height = estimate_text_height(text, font_pt=theme.BODY_PT,
                                      width=ctx.layout.content_width)
        box, tf = textbox(ctx, height)
        tf.margin_left = Pt(12)
        set_line_spacing(tf.paragraphs[0], theme.BODY_PT)
        apply_runs(tf.paragraphs[0], parse_markdown(text)[0].runs if text else [Run("")],
                    size_pt=theme.BODY_PT, color=theme.DK2)
        for run in tf.paragraphs[0].runs:
            run.font.italic = True
            run.font.color.rgb = theme.DK2
        top = ctx.cursor.place(height)
        bar = _shape(ctx, MSO_SHAPE.RECTANGLE, ctx.layout.content_left, top,
                     Pt(4), height)
        solid(bar, theme.LT2)
        no_line(bar)
        ctx.cursor.advance(GAP)

    def to_latex(self, ctx):
        from ..renderers.latex.blocks import footnotes
        from ..renderers.latex.text import blocks as md_blocks
        body = md_blocks(self.content, refs=ctx.resolve_ref) + footnotes(self, ctx)
        return "\\begin{quote}\n\\small %s\n\\end{quote}" % body

    def to_transcript(self, embedder):
        from ..renderers.transcript.text import markdown
        return f'<div class="tx-aside">{markdown(self.content)}</div>'
