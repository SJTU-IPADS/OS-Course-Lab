"""The `code` element."""
from __future__ import annotations

from ..model import Block, element
from .. import model


@element
class Code(Block):
    """A code listing, verbatim."""
    KIND = "code"
    TARGETS = frozenset({"viewer", "pptx", "latex", "transcript"})
    VERBATIM = True

    def check_content(self, page_id: str) -> None:
        """A marked code line names a line that exists, and is not blank.

        Line numbers are 1-based over the block as the slide shows it — every
        renderer strips the leading and trailing blank lines a triple-quoted
        literal collects, so the numbers have to count what is left. A number past
        the end, or one landing on a blank line, is refused rather than ignored: an
        author who renumbers a listing and forgets the mark would otherwise find out
        on the projector, from a wash that quietly went missing.
        """
        content = self.content
        if "mark" not in content:
            return
        language = content.get("language")
        if language != "pseudo":
            raise model.ValidationError(
                f"Code mark on {page_id} needs language 'pseudo', got {language!r} — "
                "another language is coloured by the renderer's own highlighter, "
                "which leaves nowhere to put the wash"
            )
        tone = content.get("tone")
        if tone not in model.HIGHLIGHT_TONES:
            raise model.ValidationError(
                f"Code mark tone on {page_id} must be one of "
                f"{sorted(model.HIGHLIGHT_TONES)}, got {tone!r}"
            )
        lines = str(content.get("content", "")).strip("\n").split("\n")
        for number in content["mark"]:
            # bool is an int subclass; reject it so True/False can't pass as a line.
            if type(number) is not int or number < 1:
                raise model.ValidationError(
                    f"Code mark on {page_id} must be a 1-based line number, "
                    f"got {number!r}"
                )
            if number > len(lines):
                raise model.ValidationError(
                    f"Code mark on {page_id} names line {number}, but the block has "
                    f"{len(lines)} line{'' if len(lines) == 1 else 's'}"
                )
            if not lines[number - 1].strip():
                raise model.ValidationError(
                    f"Code mark on {page_id} names line {number}, which is blank"
                )

    def to_marp(self, ctx):
        from .. import model
        from ..renderers.viewer.blocks import pseudo_html
        language = self.content["language"]
        content = self.content["content"].strip("\n")
        if language == "pseudo":
            marked = set(self.content.get("mark") or ())
            tone = self.content.get("tone") or model.DEFAULT_HIGHLIGHT_TONE
            return [pseudo_html(content, marked, tone), ""]
        return [f"```{language}", content, "```", ""]

    def footnote_marker(self, lines, numbers):
        from ..renderers.viewer import blocks
        return blocks.mark_tucked(lines, numbers)

    def to_pptx(self, ctx):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Emu, Pt
        from .. import model
        from ..renderers.pptx import theme
        from ..renderers.pptx.blocks import (
            GAP,
            _shape,
            no_line,
            set_line_spacing,
            solid,
            style_run,
            textbox,
        )
        from ..renderers.pptx.layout import LINE_HEIGHT
        code = str(self.content["content"]).strip("\n")
        lines = code.split("\n")
        # A marked line takes PowerPoint's own text highlight — the same wash a
        # `<mark>` gets here, at line scale.
        marked = set(self.content.get("mark") or ())
        tone = self.content.get("tone") or model.DEFAULT_HIGHLIGHT_TONE
        # Match the background to the actual code rows.  The previous ``n + 1``
        # estimate effectively reserved a blank line in every block, so short
        # listings looked especially loose.  Give each row the same exact leading
        # used elsewhere, plus a small and explicit vertical inset.
        vertical_pad_pt = 4
        height = Emu(Pt(
            len(lines) * theme.CODE_PT * LINE_HEIGHT + 2 * vertical_pad_pt
        ))
        box, tf = textbox(ctx, height)
        solid(box, theme.LT2)
        no_line(box)
        tf.margin_left = Pt(10)
        tf.margin_top = tf.margin_bottom = Pt(vertical_pad_pt)
        first = True
        for number, line in enumerate(lines, start=1):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            set_line_spacing(p, theme.CODE_PT)
            style_run(p.add_run(), line or " ", size_pt=theme.CODE_PT,
                       color=theme.FG, mono=True,
                       mark=tone if number in marked else None)
        top = ctx.cursor.place(height)
        # accent1 left bar (pre border-left).
        bar = _shape(ctx, MSO_SHAPE.RECTANGLE, ctx.layout.content_left, top,
                     Pt(4), height)
        solid(bar, theme.ACCENT1)
        no_line(bar)
        ctx.cursor.advance(GAP)

    def to_latex(self, ctx):
        from .. import model
        from ..renderers.latex.blocks import _pseudo, footnotes
        language = self.content.get("language") or ""
        if language == "pseudo":
            return _pseudo(
                self.content["content"].strip("\n"),
                set(self.content.get("mark") or ()),
                self.content.get("tone") or model.DEFAULT_HIGHLIGHT_TONE,
            ) + footnotes(self, ctx)
        option = f"[language={language}]" if language else ""
        return "\\begin{lstlisting}%s\n%s\n\\end{lstlisting}%s" % (
            option, self.content["content"], footnotes(self, ctx)
        )

    def to_transcript(self, embedder):
        from .. import model, pseudo
        from ..renderers.transcript.blocks import PSEUDO_CLASSES
        from ..renderers.transcript.text import escape
        language = self.content["language"]
        content = self.content["content"].strip("\n")
        if language != "pseudo":
            return f'<pre class="tx-code">{escape(content)}</pre>'
        marked = set(self.content.get("mark") or ())
        tone = self.content.get("tone") or model.DEFAULT_HIGHLIGHT_TONE
        lines = []
        for number, tokens in enumerate(pseudo.tokenize(content), start=1):
            line = "".join(
                f'<span class="tok-{PSEUDO_CLASSES[kind]}">{escape(text)}</span>'
                if kind in PSEUDO_CLASSES else escape(text)
                for kind, text in tokens
            )
            if number in marked:
                line = f'<span class="tx-codemark tx-codemark--{tone}">{line}</span>'
            lines.append(line)
        return '<pre class="tx-code tx-pseudo">%s</pre>' % "\n".join(lines)
