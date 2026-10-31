"""The `table` element."""
from __future__ import annotations

from ..model import Block, element


@element
class Table(Block):
    """A table; its cells are markdown and may carry file links."""
    KIND = "table"
    TARGETS = frozenset({"viewer", "pptx", "latex", "transcript"})
    FILE_LINKS = True
    CENTERED = True

    def to_marp(self, ctx):
        from ..renderers.viewer.blocks import ALIGN_DELIMITERS, fixed_table, table_cell
        table = self.content
        headers = table["headers"]
        align = table["align"]
        if align is None:
            delimiters = ["---"] * len(headers)
        else:
            delimiters = [ALIGN_DELIMITERS[a] for a in align]
        lines = [
            "| " + " | ".join(table_cell(h) for h in headers) + " |",
            "| " + " | ".join(delimiters) + " |",
        ]
        for row in table["rows"]:
            lines.append("| " + " | ".join(table_cell(c) for c in row) + " |")
        widths = table.get("widths")
        if widths is None:
            return [*lines, ""]
        return fixed_table(lines, widths)

    def footnote_marker(self, lines, numbers):
        from ..renderers.viewer import blocks
        return blocks.mark_tucked(lines, numbers)

    def to_pptx(self, ctx):
        from pptx.util import Emu, Pt
        from ..renderers.pptx import theme
        from ..renderers.pptx.blocks import GAP, fill_cell
        content = self.content
        headers = content["headers"]
        rows = content["rows"]
        align = content["align"]
        n_rows = len(rows) + 1
        n_cols = len(headers)
        row_h = Emu(Pt(theme.BODY_PT * 0.85 * 1.6))
        height = row_h * n_rows
        frame = ctx.slide.shapes.add_table(
            n_rows, n_cols, ctx.layout.content_left, ctx.cursor.top,
            ctx.layout.content_width, height,
        )
        table = frame.table
        table.first_row = True
        widths = content.get("widths")
        if widths is not None:
            # the last column takes the rounding remainder, so the sum is exact
            total = ctx.layout.content_width
            spans = [Emu(int(total * w)) for w in widths[:-1]]
            spans.append(Emu(total - sum(spans)))
            for column, span in zip(table.columns, spans):
                column.width = span
        for c, header in enumerate(headers):
            fill_cell(table.cell(0, c), str(header), bold=True, align=align, col=c)
        for r, row in enumerate(rows, start=1):
            for c, cell in enumerate(row):
                fill_cell(table.cell(r, c), str(cell), bold=False, align=align, col=c)
        ctx.cursor.place(height)
        ctx.cursor.advance(GAP)

    def to_latex(self, ctx):
        from ..renderers.latex.blocks import COLUMN_SPEC, footnotes, width_column
        from ..renderers.latex.text import inline
        headers = self.content["headers"]
        align = self.content.get("align") or ["left"] * len(headers)
        widths = self.content.get("widths")
        if widths is None:
            spec = "".join(COLUMN_SPEC[a] for a in align)
        else:
            spec = "".join(width_column(a, w) for a, w in zip(align, widths))
        lines = [
            r"\begin{center}",
            r"\begin{tabular}{%s}" % spec,
            r"\toprule",
            " & ".join(inline(h) for h in headers) + r" \\",
            r"\midrule",
        ]
        for row in self.content["rows"]:
            lines.append(" & ".join(inline(cell) for cell in row) + r" \\")
        lines += [r"\bottomrule", r"\end{tabular}", r"\end{center}"]
        return "\n".join(lines) + footnotes(self, ctx)

    def to_transcript(self, embedder):
        from ..renderers.transcript.text import inline
        content = self.content
        align = content.get("align") or ["left"] * len(content["headers"])
        head = "".join(
            f'<th style="text-align:{a}">{inline(cell)}</th>'
            for cell, a in zip(content["headers"], align)
        )
        body = "".join(
            "<tr>%s</tr>" % "".join(
                f'<td style="text-align:{a}">{inline(cell)}</td>'
                for cell, a in zip(row, align)
            )
            for row in content["rows"]
        )
        widths = content.get("widths")
        if widths is None:
            return f'<table class="tx-table"><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>'
        cols = "".join(f'<col style="width:{100 * w:.2f}%">' for w in widths)
        return (
            '<table class="tx-table" style="table-layout:fixed;width:100%">'
            f"<colgroup>{cols}</colgroup>"
            f"<thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"
        )
