"""The `demo` element."""
from __future__ import annotations

from ..model import Block, element
from .. import model


@element
class Demo(Block):
    """A command and its recorded output."""
    KIND = "demo"
    TARGETS = frozenset({"viewer", "pptx", "latex"})
    LINK_FREE_KEYS = frozenset({'output'})

    def check_content(self, page_id: str) -> None:
        """A bold output line names a line of the recorded output that exists.

        Numbered the way `Code.check_content` numbers a marked line — 1-based, over the
        output with its leading and trailing blank lines stripped — and refused on
        the same grounds: a number that points nowhere would otherwise surface on
        the projector as a row that quietly stopped being bold.
        """
        content = self.content
        shown = content.get("show")
        if shown is not None:
            if not isinstance(shown, str):
                raise model.ValidationError(f"Demo show on {page_id} must be a string")
            if len(shown.split("\n")) != len(str(content["command"]).split("\n")):
                raise model.ValidationError(
                    f"Demo show on {page_id} must have as many lines as the command: "
                    "the drawer writes each line of the command as the line of show "
                    "that stands in its place"
                )
        if "bold" not in content:
            return
        output = content.get("output")
        if not output or not str(output).strip():
            raise model.ValidationError(f"Demo bold on {page_id} needs an output= to bold")
        lines = str(output).strip("\n").split("\n")
        for number in content["bold"]:
            # bool is an int subclass; reject it so True/False can't pass as a line.
            if type(number) is not int or number < 1:
                raise model.ValidationError(
                    f"Demo bold on {page_id} must be a 1-based line number, "
                    f"got {number!r}"
                )
            if number > len(lines):
                raise model.ValidationError(
                    f"Demo bold on {page_id} names line {number}, but the output has "
                    f"{len(lines)} line{'' if len(lines) == 1 else 's'}"
                )
            if not lines[number - 1].strip():
                raise model.ValidationError(
                    f"Demo bold on {page_id} names line {number}, which is blank"
                )

    def to_marp(self, ctx):
        """A command, drawn as the transcript it is, with a button that runs it.

        The button carries the command's id, never the command — see
        ``lecturekit.demo`` for why. It ships ``disabled``: a demo is only runnable
        against a live dev server, so the static deck shows the same block as an
        inert transcript and ``demo.js`` arms it when a server is actually behind
        the page.

        Two shapes, because two things get written as demos. A bare one-liner stays
        a chip on one row, the way it always did; a command with several lines or a
        recorded output becomes a bordered transcript, so that a page which used to
        hold a ``p.code(...)`` of the same thing keeps its proportions.

        ``files=`` adds one button per named file beside the run button. They carry
        an id and not a path, and ship ``disabled``, for the same reasons the run
        button does — see ``lecturekit.source``.
        """
        from ..renderers.viewer.blocks import demo_files, demo_output, escape_html, inline_md
        from lecturekit import demo as demo_module

        content = self.content
        command = str(content["command"])
        output = content.get("output")
        lines = demo_module.prompt_lines(demo_module.shown(content))
        block_form = len(lines) > 1 or bool(output)
        # A blank command is a file button with nothing to run: no play button and
        # no `$ …` line. The file buttons are the whole of the block.
        run = ""
        if command.strip():
            run = (
                '<button class="lk-demo-run" type="button" disabled'
                ' aria-label="Run this demo">▶</button>'
            )
        button = run + demo_files(content.get("files"))
        name = f'<span class="lk-demo-name">{escape_html(content["name"])}</span>'
        desc = ""
        if content.get("description"):
            desc = f'<span class="lk-demo-desc">{inline_md(content["description"])}</span>'
        # `lk-demo-cmd` holds the command lines and nothing else — `demo.js` falls
        # back on its first line to label a run whose demo has no name.
        cmd = ""
        if lines:
            cmd = f'<code class="lk-demo-cmd">{escape_html(chr(10).join(lines))}</code>'
        open_tag = (
            f'<div class="lk-demo" data-lk-demo="{demo_module.demo_id(command, content.get("pwd"))}"'
            f' data-lk-demo-form="{"block" if block_form else "inline"}">'
        )
        if not block_form:
            return ["".join([open_tag, button, name, cmd, desc, "</div>"]), ""]
        parts = [open_tag, f'<div class="lk-demo-head">{button}{name}{desc}</div>',
                 f'<pre class="lk-demo-term">{cmd}']
        if output:
            # No newline between the two: inside a `pre` that would be a blank row,
            # and a slide cannot spare one. The gap is CSS's to draw.
            bold = content.get("bold")
            # The attribute lets the theme draw the lines `bold=` leaves out lighter,
            # so the named ones stand out by contrast from both sides.
            attr = " data-lk-demo-bold" if bold else ""
            out = demo_output(str(output), bold)
            parts.append(f'<code class="lk-demo-out"{attr}>{out}</code>')
        parts.append("</pre></div>")
        # The line breaks are written as character references, so the transcript is
        # one source line. A raw HTML block ends at the first blank line of the
        # Markdown, and a blank row in the output or in a quoted argument would be
        # one: Marp would parse the rows after it as a paragraph inside the `pre`.
        return ["".join(parts).replace("\n", "&#10;"), ""]

    def to_pptx(self, ctx):
        """A `p.demo(...)` as a still transcript.

        Nothing here is runnable — a PPTX is a file on someone else's laptop — so
        the block prints what the deck prints when nobody presses the button: the
        name, then the command, then the output the author recorded. The box is the
        code block's, because on the slide it *is* the code block, and a handout
        that reflows every listing is a different handout.
        """
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Emu, Pt
        from ..renderers.pptx import theme
        from ..renderers.pptx.blocks import (
            GAP,
            _shape,
            apply_runs,
            no_line,
            set_line_spacing,
            solid,
            style_run,
            textbox,
        )
        from ..renderers.pptx.layout import estimate_text_height, LINE_HEIGHT
        from ..renderers.pptx.text import parse_markdown
        from lecturekit import demo as demo_module

        content = self.content
        lines = demo_module.prompt_lines(demo_module.shown(content))
        output = content.get("output")
        outputs = str(output).strip("\n").split("\n") if output else []

        # A demo may go without a name, and then the listing is all there is.
        head = " — ".join(
            str(part) for part in (content["name"], content.get("description")) if part
        )
        if head:
            head_height = estimate_text_height(head, font_pt=theme.CAPTION_PT,
                                               width=ctx.layout.content_width)
            box, tf = textbox(ctx, head_height)
            set_line_spacing(tf.paragraphs[0], theme.CAPTION_PT)
            apply_runs(tf.paragraphs[0], parse_markdown(head)[0].runs,
                        size_pt=theme.CAPTION_PT, color=theme.DK2)
            for run in tf.paragraphs[0].runs:
                run.font.bold = True
            ctx.cursor.place(head_height)

        vertical_pad_pt = 4
        rows = lines + outputs
        if not rows:
            return
        bold = {len(lines) + number for number in content.get("bold") or ()}
        height = Emu(Pt(len(rows) * theme.CODE_PT * LINE_HEIGHT + 2 * vertical_pad_pt))
        box, tf = textbox(ctx, height)
        solid(box, theme.LT2)
        no_line(box)
        tf.margin_left = Pt(10)
        tf.margin_top = tf.margin_bottom = Pt(vertical_pad_pt)
        first = True
        for number, line in enumerate(rows, start=1):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            set_line_spacing(p, theme.CODE_PT)
            style_run(p.add_run(), line or " ", size_pt=theme.CODE_PT, mono=True,
                       color=theme.ACCENT1 if number <= len(lines) else theme.FG,
                       bold=number in bold)
        top = ctx.cursor.place(height)
        bar = _shape(ctx, MSO_SHAPE.RECTANGLE, ctx.layout.content_left, top,
                     Pt(4), height)
        solid(bar, theme.ACCENT4)
        no_line(bar)
        ctx.cursor.advance(GAP)

    def to_latex(self, ctx):
        """A command as the book has to print it: verbatim, with its output.

        The book cannot run anything, so the recorded output is the whole of what a
        reader gets — and it goes in a listing rather than through `inline`, because
        a shell transcript is text a machine wrote and every character in it is
        load-bearing.
        """
        from .. import i18n
        from ..renderers.latex.blocks import WEIGHT_DELIMS, footnotes
        from ..renderers.latex.text import inline
        from lecturekit import demo as demo_module

        content = self.content
        lines = [
            r"\begin{tcolorbox}[colback=blue!3,colframe=blue!40,title=%s]"
            % i18n.ui(ctx.lang, "demo")
        ]
        if content["name"]:  # a demo may go without a name
            lines.append(r"\textbf{%s}\\" % inline(content["name"]))
        if content.get("description"):
            lines.append(inline(content["description"]) + r"\\")
        transcript = demo_module.prompt_lines(demo_module.shown(content))
        options = ""
        if not transcript and not content.get("output"):
            lines.append(r"\end{tcolorbox}")
            return "\n".join(lines) + footnotes(self, ctx)
        if content.get("output"):
            outputs = str(content["output"]).strip("\n").split("\n")
            bold = set(content.get("bold") or ())
            if bold:
                # `listings` sets what sits between two invisible delimiters in the
                # style the pair names and prints everything else verbatim. The named
                # lines go bold and the rest light, as on the slide; each style takes
                # a pair the output lacks.
                (bl, br), (ll, lr) = [pair for pair in WEIGHT_DELIMS
                                      if not any(d in content["output"] for d in pair)][:2]
                options = ("[moredelim={[is][\\bfseries]{%s}{%s}},"
                           "moredelim={[is][\\fontseries{l}\\selectfont]{%s}{%s}}]"
                           % (bl, br, ll, lr))
                outputs = [f"{bl}{line}{br}" if number in bold
                           else f"{ll}{line}{lr}" if line else line
                           for number, line in enumerate(outputs, start=1)]
            transcript += outputs
        lines.append(
            "\\begin{lstlisting}%s\n%s\n\\end{lstlisting}" % (options, "\n".join(transcript))
        )
        lines.append(r"\end{tcolorbox}")
        return "\n".join(lines) + footnotes(self, ctx)
